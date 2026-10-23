"""Turns outreach.json into MSB_Outreach.xlsx for the Power Automate send flow.
Usage: python3 scripts/make_send_sheet.py outreach.json MSB_Outreach.xlsx
"""
import html, json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation

src, dst = sys.argv[1], sys.argv[2]
rows = json.load(open(src))
order = {"high": 0, "medium": 1, "unverified": 2, "low": 3}
rows.sort(key=lambda r: (r["priority"], order.get(r["confidence"], 4), r["sponsor"].lower()))

def to_html(body: str) -> str:
    paras = body.split("\n\n")
    out = []
    for p in paras:
        out.append("<p>" + "<br>".join(html.escape(line) for line in p.split("\n")) + "</p>")
    return "".join(out)

F = "Arial"
wb = Workbook()
ws = wb.active
ws.title = "Outreach"
headers = ["Send", "Status", "Sponsor", "To", "Contact", "Contact check", "Ask", "Subject", "BodyHTML", "Body", "Notes", "Slug", "Priority"]
ws.append(headers)
for r in rows:
    check = "Double-check before sending" if r["confidence"] in ("low", "unverified") else ("Good" if r["confidence"] == "high" else "Probably right")
    ws.append(["", "", r["sponsor"], r["to"], r["contact"], check, r["ask"], r["subject"], to_html(r["body"]), r["body"], r["notes"], r["slug"], r["priority"]])

n = len(rows) + 1
tbl = Table(displayName="Outreach", ref=f"A1:M{n}")
tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tbl)

widths = {"A": 8, "B": 12, "C": 34, "D": 32, "E": 24, "F": 22, "G": 40, "H": 50, "I": 14, "J": 90, "K": 50, "L": 30, "M": 9}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
yellow = PatternFill("solid", fgColor="FFFF00")
for row in ws.iter_rows(min_row=1, max_row=n):
    for c in row:
        c.font = Font(name=F, bold=(c.row == 1), size=10)
        c.alignment = Alignment(vertical="top", wrap_text=c.column_letter in ("J", "K", "G"))
for i in range(2, n + 1):
    ws[f"A{i}"].fill = yellow
    if ws[f"F{i}"].value == "Double-check before sending":
        ws[f"F{i}"].font = Font(name=F, size=10, bold=True, color="C00000")
ws.freeze_panes = "D2"
dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"A2:A{n}")

g = wb.create_sheet("How to use", 0)
lines = [
    ("MSB sponsorship outreach: send file", True),
    ("", False),
    ("1. Read each email in the Body column (Outreach tab). Edit the Subject or Body if you want, but keep BodyHTML in sync or edit BodyHTML too.", False),
    ("2. Type Yes in the yellow Send column for every email you approve. Leave it blank or No to skip.", False),
    ("3. Rows marked 'Double-check before sending' have a contact that may be out of date. Confirm it first.", False),
    ("4. Run the Power Automate flow. It sends only rows where Send = Yes and Status is blank, then writes Sent into Status.", False),
    ("5. Never rename the table (Outreach) or the column headers. The flow looks them up by name.", False),
    ("", False),
    ("Emails approved to send", True),
    ("=COUNTIF(Outreach!A2:A1000,\"Yes\")", False),
    ("Emails already sent", True),
    ("=COUNTIF(Outreach!B2:B1000,\"Sent\")", False),
    ("", False),
    ("Sender: Ishpreet Bal, Director of Finance and Scholarships (from the club's details). Club numbers used: over 100 at last year's showcase, 20 members, almost 2,000 Instagram followers (from Ishpreet).", False),
    ("Asks are sized per sponsor: $50 to $100 or in-kind for small local businesses, $100 to $250 for mid-size local businesses, $250 to $500 for large companies. McMaster funds and grants get a question email instead of an ask.", False),
]
for i, (t, bold) in enumerate(lines, start=1):
    c = g.cell(row=i, column=1, value=t)
    c.font = Font(name=F, bold=bold, size=12 if i == 1 else 10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
g["A10"].font = Font(name=F, size=10, color="000000")
g.column_dimensions["A"].width = 120
wb.save(dst)
print(f"{len(rows)} rows -> {dst}")
